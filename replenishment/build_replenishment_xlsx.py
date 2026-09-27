"""在庫0・直近7日売上ありSKUの補充リストExcelを生成する。
使い方: python3 build_replenishment_xlsx.py 在庫.csv 売上1.csv [売上2.csv ...] -o 出力.xlsx
"""
import argparse, csv
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation

import os
STOCK_MAX = int(os.environ.get("STOCK_MAX", 50000))   # 在庫データの数式を用意しておく行数
LIST_MAX = int(os.environ.get("LIST_MAX", 2000))     # 補充リストの最大表示件数
ORDER_MAX = int(os.environ.get("ORDER_MAX", 30000))   # 売上データの計算列を用意しておく行数
F = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F4E78")
HELP_FILL = PatternFill("solid", fgColor="E2EFDA")
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def read_csv(path):
    with open(path, encoding="cp932", newline="") as f:
        rows = list(csv.reader(f))
    return rows[0], rows[1:]


def num(v):
    try:
        return int(v)
    except ValueError:
        return v


def header(ws, row, labels, fill=HDR_FILL):
    for c, v in enumerate(labels, 1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.font = Font(name=F, bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BOX


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("stock")
    ap.add_argument("orders", nargs="+")
    ap.add_argument("-o", "--out", required=True)
    a = ap.parse_args()

    wb = Workbook()
    wl = wb.active
    wl.title = "補充リスト"
    ws = wb.create_sheet("在庫データ")
    wo = wb.create_sheet("売上データ")
    wh = wb.create_sheet("使い方")

    def stk(c): return f"在庫データ!${c}$2:${c}${STOCK_MAX+1}"
    def odr(c): return f"売上データ!${c}$2:${c}${ORDER_MAX+1}"

    # ---------- 売上データ（差分追記） + 計算列 ----------
    oh, _ = read_csv(a.orders[0])
    header(wo, 1, oh)
    r = 2
    for p in a.orders:
        h, rows = read_csv(p)
        assert h == oh, f"列構成が異なります: {p}"
        for row in rows:
            for c, v in enumerate(row, 1):
                if c == 2:
                    v = datetime.strptime(v, "%Y/%m/%d %H:%M:%S")
                elif c in (1, 16):          # 注文番号・JANは文字列のまま
                    pass
                else:
                    v = num(v)
                cell = wo.cell(row=r, column=c, value=v)
                cell.font = Font(name=F)
                if c == 2:
                    cell.number_format = "yyyy/mm/dd hh:mm:ss"
            r += 1
    n_orders = r - 2
    assert n_orders < ORDER_MAX
    ohelp = ["直近7日\n(1=期間内)", "SKUキー", "在庫データ\n行番号", "補充対象\n(1=対象)", "対象連番"]
    for j, v in enumerate(ohelp):
        cell = wo.cell(row=1, column=19 + j, value=v)   # S〜W列
        cell.font = Font(name=F, bold=True)
        cell.fill = HELP_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = BOX
    wo["R1"] = "←A〜Q列に追記"
    wo["R1"].font = Font(name=F, color="C00000", bold=True)
    for i in range(2, ORDER_MAX + 2):
        wo[f"S{i}"] = (f'=IF($A{i}="",0,IF(AND($B{i}>=補充リスト!$C$4,'
                       f'$B{i}<補充リスト!$C$3+1),1,0))')
        wo[f"T{i}"] = f'=IF($S{i}=1,$I{i}&"|"&$J{i}&"|"&$K{i},"")'
        wo[f"U{i}"] = f'=IF($S{i}=1,IFERROR(MATCH($T{i},{stk("N")},0),0),0)'
        # 在庫が閾値以下 かつ 期間内で同一SKUの初出行のみ1
        wo[f"V{i}"] = (f'=IF($U{i}=0,0,IF(INDEX({stk("I")},$U{i})<=補充リスト!$C$5,'
                       f'IF(COUNTIF($T$2:$T{i},$T{i})=1,1,0),0))')
        wo[f"W{i}"] = f"=N(W{i-1})+V{i}"
    wo.freeze_panes = "C2"
    wo.auto_filter.ref = f"A1:W{n_orders+1}"
    for col, w in zip("ABCDEFGHIJKLMNOPQRSTUVW",
                      [14, 20, 9, 8, 8, 8, 22, 14, 11, 9, 9, 9, 7, 9, 9, 16, 9, 14, 10, 18, 11, 10, 9]):
        wo.column_dimensions[col].width = w
    wo.row_dimensions[1].height = 32

    # ---------- 在庫データ（全件上書き） + 計算列 ----------
    sh, srows = read_csv(a.stock)
    header(ws, 1, sh)
    for i, row in enumerate(srows, 2):
        for c, v in enumerate(row, 1):
            if c in (4, 7, 9, 12):   # 属性1コード, 属性2コード, 在庫数量, 定価
                v = num(v)
            ws.cell(row=i, column=c, value=v).font = Font(name=F)
    cell = ws.cell(row=1, column=14, value="SKUキー\n(自動計算)")
    cell.font = Font(name=F, bold=True)
    cell.fill = HELP_FILL
    cell.alignment = Alignment(horizontal="center", wrap_text=True)
    cell.border = BOX
    ws["M1"] = "←A〜L列に貼付"
    ws["M1"].font = Font(name=F, color="C00000", bold=True)
    for i in range(2, STOCK_MAX + 2):
        ws[f"N{i}"] = f'=IF($A{i}="","",$A{i}&"|"&$D{i}&"|"&$G{i})'
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:N{len(srows)+1}"
    for col, w in zip("ABCDEFGHIJKLMN", [11, 30, 8, 8, 12, 8, 8, 16, 9, 16, 13, 10, 14, 18]):
        ws.column_dimensions[col].width = w
    ws.row_dimensions[1].height = 32

    # ---------- 補充リスト ----------
    wl["A1"] = "補充リスト（現在庫0以下 × 直近7日間に売上あり）"
    wl["A1"].font = Font(name=F, bold=True, size=14)
    settings = [
        (3, "基準日（集計終了日）", f"=INT(MAX({odr('B')}))", "yyyy/mm/dd",
         "初期値＝売上データの最新注文日。日付を直接入力して上書きも可。"),
        (4, "集計開始日", "=C3-6", "yyyy/mm/dd", "基準日を含む7日間（開始日 0:00 〜 基準日 23:59:59）"),
        (5, "補充対象の在庫数（以下）", 0, "0", "0 → 在庫0以下（マイナス在庫含む）を対象。変更可。"),
    ]
    for r_, label, val, fmt, note in settings:
        wl[f"B{r_}"] = label
        wl[f"B{r_}"].font = Font(name=F, bold=True)
        wl[f"C{r_}"] = val
        wl[f"C{r_}"].number_format = fmt
        wl[f"C{r_}"].font = Font(name=F, color="0000FF" if r_ != 4 else "000000")
        wl[f"C{r_}"].border = BOX
        if r_ != 4:
            wl[f"C{r_}"].fill = INPUT_FILL
        wl[f"D{r_}"] = note
        wl[f"D{r_}"].font = Font(name=F, color="595959", size=9)
    first, last = 9, 9 + LIST_MAX - 1
    summary = [
        (3, "対象SKU数", f"=MAX({odr('W')})"),
        (4, "対象品番数", f"=SUM($M${first}:$M${last})"),
        (5, "直近7日売上点数 合計(対象)", f"=SUM($I${first}:$I${last})"),
        (6, "在庫データに無い売上行数", f'=COUNTIFS({odr("S")},1,{odr("U")},0)'),
    ]
    for r_, label, f in summary:
        wl[f"H{r_}"] = label
        wl[f"H{r_}"].font = Font(name=F, bold=True)
        wl[f"J{r_}"] = f
        wl[f"J{r_}"].font = Font(name=F, bold=True)
        wl[f"J{r_}"].number_format = "#,##0"
        wl[f"J{r_}"].border = BOX
    wl["K6"] = "← 0でない場合、売上SKUが在庫データに存在しません"
    wl["K6"].font = Font(name=F, color="595959", size=9)
    wl["A7"] = f'=IF(J3>{LIST_MAX},"※対象が{LIST_MAX}件を超えています。表示しきれない行があります","")'
    wl["A7"].font = Font(name=F, color="C00000", bold=True)

    cols = ["No", "品番", "商品名", "サイズ", "カラー", "SKUコード", "JANコード",
            "現在庫数", "直近7日\n売上点数", "定価\n(税別)", "品番計\n在庫数", "品番計\n7日売上", "品番\n初出",
            "在庫行\n(参照用)"]
    header(wl, 8, cols)
    wl.row_dimensions[8].height = 32
    src = {"B": "A", "C": "B", "D": "E", "E": "H", "F": "K", "G": "J", "H": "I", "J": "L"}
    for i in range(first, last + 1):
        k = i - first + 1
        wl[f"A{i}"] = f'=IF({k}<=$J$3,{k},"")'
        wl[f"N{i}"] = f'=IF($A{i}="","",INDEX({odr("U")},MATCH($A{i},{odr("W")},0)))'
        for dst, s_ in src.items():
            wl[f"{dst}{i}"] = f'=IF($A{i}="","",INDEX({stk(s_)},$N{i}))'
        wl[f"I{i}"] = f'=IF($A{i}="","",SUMIF({odr("T")},INDEX({stk("N")},$N{i}),{odr("M")}))'
        wl[f"K{i}"] = f'=IF($A{i}="","",SUMIF({stk("A")},$B{i},{stk("I")}))'
        wl[f"L{i}"] = f'=IF($A{i}="","",SUMIFS({odr("M")},{odr("I")},$B{i},{odr("S")},1))'
        wl[f"M{i}"] = f'=IF($A{i}="","",IF(COUNTIF($B${first}:$B{i},$B{i})=1,1,0))'
        for c in "ABCDEFGHIJKLMN":
            cell = wl[f"{c}{i}"]
            cell.font = Font(name=F, color="808080") if c == "N" else Font(name=F)
            if c in "HIJKL":
                cell.number_format = "#,##0"
    wl.freeze_panes = "C9"
    wl.auto_filter.ref = f"A8:N{last}"
    for col, w in zip("ABCDEFGHIJKLMN", [6, 11, 32, 12, 16, 13, 16, 9, 9, 9, 9, 9, 6, 8]):
        wl.column_dimensions[col].width = w
    wl["H8"].comment = Comment("在庫データの在庫数量（SKU単位）", "補充リスト")
    wl["K8"].comment = Comment("同じ品番の全SKU在庫合計。0以下なら品番ごと欠品", "補充リスト")
    wl["N8"].comment = Comment("在庫データシート上の行位置（見出し除く）。計算用", "補充リスト")

    # ---------- 使い方 ----------
    guide = [
        "■ このファイルの目的",
        "現在庫が0以下のSKUのうち、直近7日間に売上があるものだけを「補充リスト」に自動抽出します。",
        "",
        "■ 更新手順（毎週 / 毎日）",
        "1. 在庫データ（週次・全件上書き）",
        "   ・「在庫データ」シートの A2:L の既存データを削除（A〜L列の2行目以降を選択 → Delete）",
        "   ・在庫CSV（stock_*.csv）をExcelで開き、見出しを除く A〜L列 のデータをコピーし、A2 に貼り付け",
        f"   ・N列（緑色見出し）は自動計算列です。削除・上書きしないでください（{STOCK_MAX:,}行分用意済み）",
        "2. 売上データ（週次 or 日次・差分追記）",
        "   ・売上CSV（order_*.csv）をExcelで開き、見出しを除く A〜Q列 をコピー",
        "   ・「売上データ」シートの最終行の次の行（空いているA列の先頭）に貼り付け",
        "   ・同じ期間のファイルを二重に貼り付けると売上点数が二重計上されます",
        f"   ・S〜W列（緑色見出し）は自動計算列です。削除・上書きしないでください（{ORDER_MAX:,}行分用意済み。超える場合は古い行を削除するか、最終行の数式を下へコピー）",
        "3. 「補充リスト」シートを開くと自動で再計算されます",
        "   ・基準日（C3）は売上データの最新注文日が自動で入ります。任意日付に上書きも可能です",
        "   ・在庫条件（C5）は 0 = 在庫0以下（マイナス在庫を含む）が対象",
        "",
        "■ 注意点",
        "・注文日時（売上データB列）が日付として認識されている必要があります（CSVをExcelで開けば自動変換されます）",
        "・在庫と売上は「商品コード＋属性１コード＋属性２コード」で照合しています（JANは在庫側で重複があるため不使用）",
        "・売上点数は「数量」列の合計です。単価0円の注文も数量として計上しています",
        f"・補充リストは最大{LIST_MAX}件まで表示します。超えた場合はリスト上部に警告が出ます",
        "・補充リストの並び順は直近7日間で最初に売れた順です。フィルタ（見出しの▼）で並べ替え・絞り込みができます",
        "・売上SKUが在庫データに存在しない場合は補充リストに出ません（件数は補充リストJ6に表示）",
    ]
    for i, t in enumerate(guide, 1):
        wh[f"A{i}"] = t
        wh[f"A{i}"].font = Font(name=F, bold=t.startswith("■"), size=12 if t.startswith("■") else 10)
    wh.column_dimensions["A"].width = 110

    wb.move_sheet("使い方", offset=-3)
    wb.active = 1
    wb.save(a.out)


if __name__ == "__main__":
    main()
