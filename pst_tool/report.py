"""Render AnalysisResult objects as JSON, Markdown, or CSV."""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Iterable

from .analyzer import AnalysisResult, MessageRecord


def _human_size(num_bytes: int) -> str:
    size = float(num_bytes)
    for unit in ("B", "KB", "MB", "GB", "TB"):
        if size < 1024 or unit == "TB":
            return f"{size:.1f} {unit}"
        size /= 1024
    return f"{size:.1f} TB"


def to_dict(result: AnalysisResult, *, top_n: int = 10) -> dict:
    return {
        "root": str(result.root),
        "total_messages": result.total_messages,
        "parse_errors": result.parse_errors,
        "total_size_bytes": result.total_size_bytes,
        "total_size_human": _human_size(result.total_size_bytes),
        "date_range": {
            "earliest": result.earliest.isoformat() if result.earliest else None,
            "latest": result.latest.isoformat() if result.latest else None,
        },
        "folders": dict(result.folder_counts.most_common()),
        "top_senders": result.top(result.sender_counts, top_n),
        "top_recipients": result.top(result.recipient_counts, top_n),
        "top_domains": result.top(result.domain_counts, top_n),
        "messages_per_month": dict(sorted(result.monthly_counts.items())),
        "attachments": {
            "total_count": result.total_attachments,
            "total_bytes": result.total_attachment_bytes,
            "total_human": _human_size(result.total_attachment_bytes),
            "by_extension": dict(result.attachment_extension_counts.most_common()),
        },
    }


def to_json(result: AnalysisResult, *, top_n: int = 10, indent: int = 2) -> str:
    return json.dumps(to_dict(result, top_n=top_n), ensure_ascii=False, indent=indent)


def to_markdown(result: AnalysisResult, *, top_n: int = 10) -> str:
    d = to_dict(result, top_n=top_n)
    lines = [
        f"# PST 解析レポート: `{d['root']}`",
        "",
        f"- 総メッセージ数: **{d['total_messages']}**",
        f"- 解析エラー: {d['parse_errors']}",
        f"- 合計サイズ: {d['total_size_human']}",
        f"- 期間: {d['date_range']['earliest'] or 'N/A'} 〜 {d['date_range']['latest'] or 'N/A'}",
        f"- 添付ファイル数: {d['attachments']['total_count']} ({d['attachments']['total_human']})",
        "",
        "## フォルダ別メッセージ数",
        "",
    ]
    for folder, count in d["folders"].items():
        lines.append(f"- {folder}: {count}")

    lines += ["", f"## 送信者トップ {top_n}", ""]
    for sender, count in d["top_senders"]:
        lines.append(f"- {sender}: {count}")

    lines += ["", f"## 受信者トップ {top_n}", ""]
    for recipient, count in d["top_recipients"]:
        lines.append(f"- {recipient}: {count}")

    lines += ["", f"## ドメイン別送信者トップ {top_n}", ""]
    for domain, count in d["top_domains"]:
        lines.append(f"- {domain}: {count}")

    lines += ["", "## 月別メッセージ数", ""]
    for month, count in d["messages_per_month"].items():
        lines.append(f"- {month}: {count}")

    lines += ["", "## 添付ファイルの拡張子別件数", ""]
    for ext, count in d["attachments"]["by_extension"].items():
        lines.append(f"- {ext}: {count}")

    return "\n".join(lines) + "\n"


CSV_FIELDS = [
    "folder",
    "date",
    "sender",
    "sender_name",
    "to",
    "cc",
    "subject",
    "size_bytes",
    "attachment_count",
    "attachment_names",
    "path",
]


def _message_row(m: MessageRecord) -> list:
    return [
        m.folder,
        m.date.isoformat() if m.date else "",
        m.sender,
        m.sender_name,
        "; ".join(m.to),
        "; ".join(m.cc),
        m.subject,
        m.size_bytes,
        len(m.attachments),
        "; ".join(a.filename for a in m.attachments),
        str(m.path),
    ]


class CsvStreamWriter:
    """Write message rows to CSV one at a time.

    Use this (instead of ``write_csv``) together with ``analyzer.analyze``'s
    ``on_message`` callback when the mailbox is too large to hold every
    ``MessageRecord`` in memory at once.
    """

    def __init__(self, out_path: Path):
        self._fh = Path(out_path).open("w", newline="", encoding="utf-8")
        self._writer = csv.writer(self._fh)
        self._writer.writerow(CSV_FIELDS)

    def write(self, m: MessageRecord) -> None:
        self._writer.writerow(_message_row(m))

    def close(self) -> None:
        self._fh.close()

    def __enter__(self) -> "CsvStreamWriter":
        return self

    def __exit__(self, *exc_info) -> None:
        self.close()


def write_csv(messages: Iterable[MessageRecord], out_path: Path) -> None:
    with CsvStreamWriter(out_path) as writer:
        for m in messages:
            writer.write(m)


class XlsxReportWriter:
    """Build a multi-sheet .xlsx report (Excel-friendly output).

    Messages are streamed into the workbook one row at a time via
    ``write()`` (uses openpyxl's write-only mode, so memory stays flat
    regardless of mailbox size). Call ``close(result)`` at the end to add
    the summary/ranking sheets (built from the already-aggregated
    ``AnalysisResult``, which is cheap even for huge mailboxes) and save
    the file.
    """

    def __init__(self, out_path: Path):
        try:
            import openpyxl
        except ImportError as exc:
            raise RuntimeError(
                "openpyxl がインストールされていません。`pip install openpyxl` を実行してください。"
            ) from exc
        self._out_path = Path(out_path)
        self._wb = openpyxl.Workbook(write_only=True)
        self._messages_ws = self._wb.create_sheet("Messages")
        self._messages_ws.append(CSV_FIELDS)

    def write(self, m: MessageRecord) -> None:
        self._messages_ws.append(_message_row(m))

    def close(self, result: AnalysisResult, *, top_n: int = 10) -> None:
        d = to_dict(result, top_n=top_n)

        summary_ws = self._wb.create_sheet("Summary")
        summary_ws.append(["項目", "値"])
        summary_ws.append(["総メッセージ数", d["total_messages"]])
        summary_ws.append(["解析エラー", d["parse_errors"]])
        summary_ws.append(["合計サイズ", d["total_size_human"]])
        summary_ws.append(["期間(開始)", d["date_range"]["earliest"] or ""])
        summary_ws.append(["期間(終了)", d["date_range"]["latest"] or ""])
        summary_ws.append(["添付ファイル数", d["attachments"]["total_count"]])
        summary_ws.append(["添付ファイル合計サイズ", d["attachments"]["total_human"]])

        folders_ws = self._wb.create_sheet("Folders")
        folders_ws.append(["フォルダ", "件数"])
        for folder, count in d["folders"].items():
            folders_ws.append([folder, count])

        senders_ws = self._wb.create_sheet("Top Senders")
        senders_ws.append(["送信者", "件数"])
        for sender, count in d["top_senders"]:
            senders_ws.append([sender, count])

        recipients_ws = self._wb.create_sheet("Top Recipients")
        recipients_ws.append(["受信者", "件数"])
        for recipient, count in d["top_recipients"]:
            recipients_ws.append([recipient, count])

        domains_ws = self._wb.create_sheet("Top Domains")
        domains_ws.append(["ドメイン", "件数"])
        for domain, count in d["top_domains"]:
            domains_ws.append([domain, count])

        monthly_ws = self._wb.create_sheet("Monthly")
        monthly_ws.append(["年月", "件数"])
        for month, count in d["messages_per_month"].items():
            monthly_ws.append([month, count])

        att_ws = self._wb.create_sheet("Attachments by Ext")
        att_ws.append(["拡張子", "件数"])
        for ext, count in d["attachments"]["by_extension"].items():
            att_ws.append([ext, count])

        self._wb.save(self._out_path)


def write_xlsx(result: AnalysisResult, out_path: Path, *, top_n: int = 10) -> None:
    """Convenience: build a full .xlsx report from an already-collected
    ``AnalysisResult`` (i.e. ``result.messages`` populated). For very large
    mailboxes, use ``XlsxReportWriter`` directly with ``analyzer.analyze``'s
    ``on_message`` streaming instead.
    """
    writer = XlsxReportWriter(out_path)
    for m in result.messages:
        writer.write(m)
    writer.close(result, top_n=top_n)
